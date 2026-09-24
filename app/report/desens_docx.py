# -*- coding: utf-8 -*-
"""把一次冻结的脱敏率结果写成 Word 报告。不引用 tkinter。"""

import os
import time
from datetime import datetime

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from app import settings as cfg


def capture_param_snapshot():
    """读当前检测参数，供报告方法节使用。必须在检测结束时调用，不能拖到导出时。"""
    snapshot = {
        "YUNET_SCORE_THRESHOLD": cfg.YUNET_SCORE_THRESHOLD,
        "MIN_FACE_SIZE": cfg.MIN_FACE_SIZE,
        "SKIN_FILTER_ENABLED": cfg.SKIN_FILTER_ENABLED,
        "PREPROCESS_RESIZE_ENABLED": cfg.PREPROCESS_RESIZE_ENABLED,
        "PREPROCESS_RESIZE_WIDTH": cfg.PREPROCESS_RESIZE_WIDTH,
        "MOSAIC_MASK_ENABLED": cfg.MOSAIC_MASK_ENABLED,
    }
    return snapshot


def write_desens_docx(report, path, is_cancelled):
    """把 report 写成 path。is_cancelled() 为真时停笔并删除本次已写出的文件。

    返回 "ok" / "cancelled"。写盘失败时删掉本次生成物后再抛出。
    """
    temp = path + ".part"
    replaced = False
    result = "ok"
    error = None
    try:
        # 还没动笔就被取消
        if is_cancelled():
            result = "cancelled"
        else:
            doc = _build_document(report, is_cancelled)
            # 拼表过程中取消：还没有目标文件
            if is_cancelled():
                result = "cancelled"
            else:
                doc.save(temp)
                # 临时文件已经落地，替换前再看一次
                if is_cancelled():
                    result = "cancelled"
                else:
                    os.replace(temp, path)
                    replaced = True
                    # 已经写到目标路径，完成提示出现前仍可取消并删掉
                    if is_cancelled():
                        result = "cancelled"
    except Exception as exc:
        # 写到一半失败：记下来，下面统一删本次生成物
        result = "fail"
        error = exc
    if result != "ok":
        _remove_generated(temp, path if replaced else None)
    if error is not None:
        raise error
    return result


def _remove_generated(temp, target):
    """删除本次临时文件；target 非空时连本次已替换的目标文件一起删。"""
    for item in (temp, target):
        # 没产生的路径跳过
        if item and os.path.exists(item):
            os.remove(item)
    return


def _build_document(report, is_cancelled):
    """按五节组装文档。取消标记在漏帧行之间检查。"""
    doc = Document()
    _set_font(doc)
    doc.add_heading("人脸脱敏情况检测报告", level=0)
    doc.add_paragraph("生成时间：%s" % datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    doc.add_heading("1. 检测对象", level=1)
    _add_subject(doc, report)

    doc.add_heading("2. 检测方法", level=1)
    _add_method(doc, report)

    doc.add_heading("3. 检测结论", level=1)
    _add_conclusion(doc, report)

    doc.add_heading("4. 漏脱敏明细", level=1)
    _add_misses(doc, report, is_cancelled)

    doc.add_heading("5. 说明", level=1)
    doc.add_paragraph("本报告按人脸张数统计脱敏率。分母是帧率换算后原视频应脱敏人脸数，分子是各次对齐后已脱敏的人脸数。原视频人脸数、脱敏视频人脸数是换算前各自检出的总数。帧数用同一套名称，不代替人脸数。")
    paint = "打开" if report["params"]["MOSAIC_MASK_ENABLED"] else "关闭"
    doc.add_paragraph("本次检测时「预处理涂抹马赛克」为%s。该开关打开时，半透明马赛克会先被涂成纯色再检人脸，脱敏视频更容易被判为检不出人脸。" % paint)
    return doc


def _set_font(doc):
    """正文使用微软雅黑，便于 Word / WPS 打开。"""
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
    return


def _add_table(doc, headers, rows, is_cancelled=None):
    """写一张表。行数多时分批让出 GIL，界面才能响应取消和任务栏。"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for col, text in enumerate(headers):
        table.rows[0].cells[col].text = text
    for r, row in enumerate(rows, start=1):
        # 每写一批看一次取消，并让出时间片给主线程
        if r % 20 == 0:
            if is_cancelled is not None and is_cancelled():
                break
            time.sleep(0.001)
        cells = table.add_row().cells
        for col, text in enumerate(row):
            cells[col].text = str(text)
    return table


def _clock(seconds):
    """秒数转时分秒。不足一小时不写小时。"""
    sec = int(round(seconds)) if seconds else 0
    if sec < 0:
        sec = 0
    hours, rem = divmod(sec, 3600)
    minutes, secs = divmod(rem, 60)
    text = "%02d:%02d" % (minutes, secs)
    # 超过一小时才带小时，避免短视频前面多一个 0:
    if hours:
        text = "%d:%02d:%02d" % (hours, minutes, secs)
    return text


def _frame_clock(frame, fps):
    """帧号对应的时间。"""
    text = "00:00"
    if fps:
        text = _clock(frame / float(fps))
    return text


def _duration(total, fps):
    """总帧对应的时长。"""
    text = "00:00"
    if fps:
        text = _clock(total / float(fps))
    return text


def _res(info):
    """宽×高；缺尺寸时写未知。"""
    text = "未知"
    if info.get("width") and info.get("height"):
        text = "%s×%s" % (info["width"], info["height"])
    return text


def _add_subject(doc, report):
    """检测对象表。"""
    src = report["src"]
    dst = report["dst"]
    rows = [
        ("文件", src["path"], dst["path"]),
        ("时长", _duration(src["total"], src["fps"]), _duration(dst["total"], dst["fps"])),
        ("分辨率", _res(src), _res(dst)),
        ("帧率", "%.2f fps" % src["fps"], "%.2f fps" % dst["fps"]),
        ("总帧", str(src["total"]), str(dst["total"])),
        ("人脸帧数", str(src["face_frames"]), str(dst["face_frames"])),
        ("人脸数", str(src.get("face_count", 0)), str(dst.get("face_count", 0))),
    ]
    _add_table(doc, ("项目", "原视频", "脱敏视频"), rows)
    return


def _on_off(value):
    """布尔显示成开/关。"""
    text = "关闭"
    if value:
        text = "打开"
    return text


def _add_method(doc, report):
    """口径、对齐和参数快照。"""
    src = report["src"]
    dst = report["dst"]
    sx, sy = report["scale"]
    offset = report["offset"]
    doc.add_paragraph("判定口径：脱敏率 = 已脱敏人脸数 ÷ 帧率换算后原视频应脱敏人脸数。每一次时间对齐，已脱敏人脸数 = 该原帧人脸数 − 该脱敏帧人脸数，不够减时记 0。不根据马赛克纹理判定。")
    doc.add_paragraph("换算公式：脱敏帧号 = round(原帧号 ÷ 原帧率 × 脱敏帧率) + 帧偏移。")
    doc.add_paragraph("帧率换算后原视频应脱敏人脸数 = 换算后落在脱敏视频时长内的原视频人脸数之和，超出时长的不计入。帧率换算后脱敏视频检测出人脸数 = 这些原帧落到的脱敏帧上检出的人脸数，同一张脱敏帧只计一次。")
    doc.add_paragraph("帧数用同一套说法：原视频人脸帧数、脱敏视频人脸帧数是各自检出过人脸的帧数；帧率换算后原视频应脱敏帧数是落在脱敏视频时长内的原视频人脸帧数；帧率换算后脱敏视频检测出帧数是上述位置上仍检出人脸的脱敏帧数，同一张脱敏帧只计一次。")
    doc.add_paragraph("坐标换算：×%.3f / ×%.3f（宽 / 高）。" % (sx, sy))
    delay = 0.0
    if dst["fps"]:
        delay = offset / float(dst["fps"])
    doc.add_paragraph("帧偏移：%+d 帧（%.2f 秒）。" % (offset, delay))
    params = report["params"]
    doc.add_paragraph("本次检测参数：")
    _add_table(doc, ("参数", "取值"), [
        ("检测置信度阈值", params["YUNET_SCORE_THRESHOLD"]),
        ("最小人脸边长(px)", params["MIN_FACE_SIZE"]),
        ("肤色过滤", _on_off(params["SKIN_FILTER_ENABLED"])),
        ("预处理缩放", _on_off(params["PREPROCESS_RESIZE_ENABLED"])),
        ("缩放目标宽度(px)", params["PREPROCESS_RESIZE_WIDTH"]),
        ("预处理涂抹马赛克", _on_off(params["MOSAIC_MASK_ENABLED"])),
    ])
    return


def _add_conclusion(doc, report):
    """结论。无人脸时不写百分比。越界数只在大于 0 时出现。"""
    # 原视频没有人脸，脱敏率不适用
    if report["na"]:
        doc.add_paragraph("原视频未检测到人脸，无需脱敏，脱敏率不适用。")
    else:
        doc.add_paragraph("换算公式：脱敏帧号 = round(原帧号 ÷ 原帧率 × 脱敏帧率) + 帧偏移。落在脱敏视频时长内的原视频人脸才计入应脱敏；超出时长的不计入。")
        _add_table(doc, ("项目", "数值"), [
            ("原视频人脸数", report["src"].get("face_count", 0)),
            ("脱敏视频人脸数", report["dst"].get("face_count", 0)),
            ("帧率换算后原视频应脱敏人脸数", report["faces_should"]),
            ("帧率换算后脱敏视频检测出人脸数", report["faces_unique_dst"]),
            ("已脱敏人脸数", report["faces_done"]),
            ("脱敏率", "%.2f%%" % report["rate"]),
            ("原视频人脸帧数", report["src"]["face_frames"]),
            ("脱敏视频人脸帧数", report["dst"]["face_frames"]),
            ("帧率换算后原视频应脱敏帧数", report["should"]),
            ("帧率换算后脱敏视频检测出帧数", report["detected_dst_frames"]),
            ("帧率换算后原视频漏脱敏帧数", report["missed"]),
        ])
        # 对齐超出脱敏视频时长的帧不进分母
        if report["skipped"]:
            doc.add_paragraph("对齐越界未计入：%d 帧。" % report["skipped"])
    return


def _add_misses(doc, report, is_cancelled):
    """全量漏帧表。没有漏帧时写明未发现。"""
    missed = [row for row in report["rows"] if not row["desensitized"]]
    # 无人脸或全部已脱敏
    if report["na"] or not missed:
        doc.add_paragraph("未发现漏脱敏帧。")
    else:
        src_fps = report["src"]["fps"]
        dst_fps = report["dst"]["fps"]
        headers = ("原视频时间", "原视频帧号", "脱敏视频时间", "脱敏视频帧号",
                   "状态", "该帧原视频人脸数", "该帧脱敏视频人脸数")
        body = []
        for index, row in enumerate(missed):
            # 每写一批看一次取消，避免长表无法停
            if index % 20 == 0 and is_cancelled():
                break
            body.append((
                _frame_clock(row["src_frame"], src_fps),
                row["src_frame"],
                _frame_clock(row["dst_frame"], dst_fps),
                row["dst_frame"],
                row["status"],
                row["src_faces"],
                row["dst_faces"],
            ))
        # 取消发生在表写完之前就不再落表；写的过程中仍可取消
        if not is_cancelled():
            _add_table(doc, headers, body, is_cancelled)
    return
