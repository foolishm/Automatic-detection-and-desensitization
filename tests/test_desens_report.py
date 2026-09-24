# -*- coding: utf-8 -*-
"""脱敏检测报告：按钮时机、docx 内容、取消不留文件。

用法：python tests/test_desens_report.py
"""

import os
import shutil
import sys
import tempfile
import tkinter as tk
from tkinter import filedialog

import _bootstrap  # noqa: F401

import face_video_detector as f


def _faces(spec):
    """(帧号, 人脸数) 列表 → 预处理 faces 记录。"""
    rows = []
    for frame, count in spec:
        boxes = [(100, 100, 40, 50)] * count
        rows.append((frame, boxes))
    return rows


def _pair(src_faces, dst_faces):
    """两路假记录。路径名用来核对默认文件名。"""
    src = {"path": "D:/样例/原视频.mp4", "fps": 10.0, "total": 30, "done": True,
           "width": 1280, "height": 720,
           "faces": _faces(src_faces), "mosaics": {}}
    dst = {"path": "D:/样例/脱敏视频.mp4", "fps": 10.0, "total": 30, "done": True,
           "width": 1280, "height": 720,
           "faces": _faces(dst_faces), "mosaics": {}}
    return src, dst


class _App(object):
    """建一个藏起来的主窗口，结束时关掉。"""

    def __init__(self):
        self.root = tk.Tk()
        self.root.withdraw()
        self.app = f.FaceVideoApp(self.root)
        self.root.update()

    def close(self):
        try:
            self.app._on_close()
        except Exception:
            self.root.destroy()
        return


def _docx_text(path):
    """把段落和表格收成一段文字，方便断言。"""
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _run_compute(app, src, dst):
    """同步跑一次脱敏率计算，并处理回到主线程的回调。"""
    app.src_info = src
    app.dst_info = dst
    app._compute_desensitization()
    app.root.update()
    return


def _export_sync(app):
    """测试里把保存线程改成当场执行，避免在没有 mainloop 时调用 root.after。"""
    import threading
    origin = threading.Thread

    class _Now(object):
        def __init__(self, target=None, daemon=None):
            self._target = target

        def start(self):
            self._target()
            return

    threading.Thread = _Now
    try:
        app.export_desens_report()
        app.root.update()
    finally:
        threading.Thread = origin
    return


def check_button_and_miss_table(tmp):
    """未检测时按钮禁用；算出 3/2 后可点，文档结论和漏帧表与界面一致。"""
    host = _App()
    ok = host.app.btn_export_report.cget("state") == tk.DISABLED
    notes = []
    saved = {}

    def fake_ask(**kwargs):
        saved["initialfile"] = kwargs.get("initialfile", "")
        return os.path.join(tmp, "out.docx")

    def fake_open(parent, on_cancel):
        notes.append("正在保存")

        class _Dlg(object):
            def close(self):
                return

        return _Dlg()

    def fake_info(title, message, parent=None, **kwargs):
        notes.append(title)
        notes.append(message)
        return

    old_ask = filedialog.asksaveasfilename
    old_open = f.messagebox.open_saving
    old_info = f.messagebox.showinfo
    path = os.path.join(tmp, "out.docx")
    try:
        filedialog.asksaveasfilename = fake_ask
        f.messagebox.open_saving = fake_open
        f.messagebox.showinfo = fake_info
        src, dst = _pair([(0, 1), (10, 2), (20, 1)], [(10, 1)])
        _run_compute(host.app, src, dst)
        text = host.app.desens_result.get("1.0", tk.END)
        ok = ok and "原视频人脸数: 4 · 脱敏视频人脸数: 1" in text
        ok = ok and "帧率换算后原视频应脱敏人脸数: 4 · 帧率换算后脱敏视频检测出人脸数: 1" in text
        ok = ok and "已脱敏人脸数: 3" in text
        ok = ok and "换算公式:" in text
        ok = ok and host.app.btn_export_report.cget("state") == tk.NORMAL
        _export_sync(host.app)
        ok = ok and saved.get("initialfile") == "原视频_vs_脱敏视频_脱敏检测报告.docx"
        ok = ok and os.path.isfile(path)
        body = _docx_text(path) if os.path.isfile(path) else ""
        ok = ok and "原视频人脸数" in body and "4" in body
        ok = ok and "帧率换算后原视频应脱敏人脸数" in body and "4" in body
        ok = ok and "帧率换算后脱敏视频检测出人脸数" in body and "1" in body
        ok = ok and "已脱敏人脸数" in body and "3" in body
        ok = ok and "该帧原视频人脸数" in body
        ok = ok and "75.00%" in body
        ok = ok and "换算公式" in body
        ok = ok and "帧率换算后原视频应脱敏帧数" in body
        ok = ok and "帧率换算后脱敏视频检测出帧数" in body
        ok = ok and "00:01" in body
        ok = ok and "未脱敏(仍检出人脸)" in body
        ok = ok and "原视频人脸数" in body and "脱敏视频人脸数" in body
        # 漏的是原帧 10，原视频 2 张脸，脱敏视频 1 张
        ok = ok and "2" in body and "1" in body
        ok = ok and "达标" not in body and "检测人" not in body
        ok = ok and notes[0] == "正在保存" and "保存完成" in notes
        ok = ok and path in "".join(notes)
    finally:
        filedialog.asksaveasfilename = old_ask
        f.messagebox.open_saving = old_open
        f.messagebox.showinfo = old_info
        host.close()
    return ok


def check_zero_miss_and_cancel_dialog(tmp):
    """没有漏帧时写明未发现；关掉另存为不产生文件。"""
    host = _App()
    ok = True
    path = os.path.join(tmp, "zero.docx")

    def fake_ask_cancel(**kwargs):
        return ""

    old_ask = filedialog.asksaveasfilename
    try:
        src, dst = _pair([(0, 1), (10, 1)], [])
        _run_compute(host.app, src, dst)
        report = host.app.last_desens_report
        outcome = f.write_desens_docx(report, path, lambda: False)
        body = _docx_text(path)
        ok = outcome == "ok" and "未发现漏脱敏帧" in body
        os.remove(path)
        filedialog.asksaveasfilename = fake_ask_cancel
        _export_sync(host.app)
        ok = ok and not os.path.exists(path)
        ok = ok and not os.path.exists(path + ".part")
    finally:
        filedialog.asksaveasfilename = old_ask
        host.close()
    return ok


def check_cancel_deletes_and_success_order(tmp):
    """替换之后取消要删掉本次文件；未覆盖的旧文件保留；成功提示在正在保存之后。"""
    report_host = _App()
    ok = True
    try:
        src, dst = _pair([(0, 1)], [])
        _run_compute(report_host.app, src, dst)
        report = report_host.app.last_desens_report
        target = os.path.join(tmp, "报告.docx")
        # 第 4 次询问发生在替换到目标路径之后
        flags = iter([False, False, False, True])

        def cancel_after_replace():
            try:
                flag = next(flags)
            except StopIteration:
                flag = True
            return flag

        outcome = f.write_desens_docx(report, target, cancel_after_replace)
        ok = outcome == "cancelled"
        ok = ok and not os.path.exists(target)
        ok = ok and not os.path.exists(target + ".part")

        old = os.path.join(tmp, "旧报告.docx")
        with open(old, "w", encoding="utf-8") as handle:
            handle.write("OLD")
        # 临时文件写完、还没替换时取消：旧文件必须还在
        early = iter([False, False, True])

        def cancel_before_replace():
            try:
                flag = next(early)
            except StopIteration:
                flag = True
            return flag

        outcome = f.write_desens_docx(report, old, cancel_before_replace)
        ok = ok and outcome == "cancelled"
        ok = ok and os.path.isfile(old)
        with open(old, "r", encoding="utf-8") as handle:
            ok = ok and handle.read() == "OLD"
        ok = ok and not os.path.exists(old + ".part")
    finally:
        report_host.close()

    host = _App()
    notes = []
    path = os.path.join(tmp, "完成.docx")

    def fake_ask(**kwargs):
        return path

    def fake_open(parent, on_cancel):
        notes.append("正在保存")

        class _Dlg(object):
            def close(self):
                return

        return _Dlg()

    def fake_info(title, message, parent=None, **kwargs):
        notes.append(title + ":" + message)
        return

    old_ask = filedialog.asksaveasfilename
    old_open = f.messagebox.open_saving
    old_info = f.messagebox.showinfo
    try:
        filedialog.asksaveasfilename = fake_ask
        f.messagebox.open_saving = fake_open
        f.messagebox.showinfo = fake_info
        src, dst = _pair([(0, 1)], [])
        _run_compute(host.app, src, dst)
        _export_sync(host.app)
        ok = ok and notes[0] == "正在保存"
        ok = ok and any(item.startswith("保存完成") and path in item for item in notes)
        ok = ok and os.path.isfile(path)
    finally:
        filedialog.asksaveasfilename = old_ask
        f.messagebox.open_saving = old_open
        f.messagebox.showinfo = old_info
        host.close()

    # 保存过程中取消：不弹完成，并且本次写出的文件被删掉
    host = _App()
    notes = []
    victim = os.path.join(tmp, "取消.docx")

    def fake_ask2(**kwargs):
        return victim

    def cancel_on_open(parent, on_cancel):
        notes.append("正在保存")
        on_cancel()

        class _Dlg(object):
            def close(self):
                return

        return _Dlg()

    def fake_info2(title, message, parent=None, **kwargs):
        notes.append(title)
        return

    # 包一层真实写入：先让它写完再看取消标记。这里用事件在 open 时已置位，
    # 另做一次「写完再取消」由上面的 write_desens_docx 覆盖。
    try:
        filedialog.asksaveasfilename = fake_ask2
        f.messagebox.open_saving = cancel_on_open
        f.messagebox.showinfo = fake_info2
        src, dst = _pair([(0, 1)], [])
        _run_compute(host.app, src, dst)
        _export_sync(host.app)
        ok = ok and "保存完成" not in notes
        ok = ok and not os.path.exists(victim)
        ok = ok and not os.path.exists(victim + ".part")
    finally:
        filedialog.asksaveasfilename = old_ask
        f.messagebox.open_saving = old_open
        f.messagebox.showinfo = old_info
        host.close()
    return ok


def main():
    f.load_params()
    tmp = tempfile.mkdtemp(prefix="desens_report_")
    ok = True
    try:
        ok = check_button_and_miss_table(tmp) and ok
        print("3.1", "PASS" if ok else "FAIL")
        step = check_zero_miss_and_cancel_dialog(tmp)
        ok = step and ok
        print("3.2", "PASS" if step else "FAIL")
        step = check_cancel_deletes_and_success_order(tmp)
        ok = step and ok
        print("3.3", "PASS" if step else "FAIL")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
